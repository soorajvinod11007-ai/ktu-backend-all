import sys
from pathlib import Path
import openpyxl
from sqlalchemy import text

sys.path.insert(0, '.')
from app.core.database import SessionLocal

def main():
    db = SessionLocal()
    print("=" * 65)
    print("           DATABASE VERIFICATION REPORT (DATA FOLDER)       ")
    print("=" * 65)

    # 1. Source Excel Check
    excel_path = Path("data/CS,ME,MC,CE.xlsx")
    if excel_path.exists():
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        print(f"\n[1] SOURCE EXCEL: {excel_path.name}")
        total_excel = 0
        for sheet in wb.sheetnames:
            rows = [r for r in wb[sheet].iter_rows(values_only=True) if any(r)]
            header_offset = 1 if len(rows) > 0 else 0
            data_rows = max(0, len(rows) - header_offset)
            total_excel += data_rows
            print(f"  - Sheet {sheet:<10}: {data_rows:>4} student rows")
        print(f"  Total students in Excel: {total_excel}")

    # 2. ZIPs & PDFs in Data folder
    data_dir = Path("data")
    zips = list(data_dir.glob("*.zip"))
    pdfs = list(data_dir.rglob("*.pdf"))
    print(f"\n[2] SOURCE ARCHIVES & FILES IN data/:")
    print(f"  - ZIP files: {len(zips)}")
    print(f"  - PDF files found across folders: {len(pdfs)}")

    # 3. DB Table Counts
    print(f"\n[3] DATABASE RECORD COUNTS:")
    tables = [
        "programs", "batches", "students", "exam_sessions", 
        "subjects", "faculty", "course_offerings", 
        "grade_cards", "semester_results"
    ]
    for t in tables:
        cnt = db.execute(text(f"SELECT COUNT(*) FROM {t}")).scalar()
        print(f"  - {t:<18}: {cnt:>5} records")

    # 4. Students breakdown
    print(f"\n[4] STUDENTS BY PROGRAM:")
    rows = db.execute(text("""
        SELECT p.code, p.name, COUNT(s.id) as cnt
        FROM programs p
        LEFT JOIN students s ON s.program_id = p.id
        GROUP BY p.code, p.name
        ORDER BY cnt DESC
    """)).fetchall()
    for r in rows:
        print(f"  - {r[0]:<5} | {r[1]:<40} : {r[2]:>3} students")

    print(f"\n[5] STUDENTS BY BATCH:")
    rows = db.execute(text("""
        SELECT b.label, b.start_year, b.end_year, COUNT(s.id) as cnt
        FROM batches b
        LEFT JOIN students s ON s.batch_id = b.id
        GROUP BY b.id
        ORDER BY b.start_year
    """)).fetchall()
    for r in rows:
        print(f"  - {r[0]:<12} ({r[1]}-{r[2]}) : {r[3]:>3} students")

    # Flagged / Auto-created students
    flagged = db.execute(text("SELECT COUNT(*) FROM students WHERE is_flagged = 1")).scalar()
    print(f"\n[6] AUTO-CREATED / FLAGGED STUDENTS: {flagged}")
    if flagged > 0:
        samples = db.execute(text("SELECT register_no, name FROM students WHERE is_flagged = 1 LIMIT 5")).fetchall()
        for s in samples:
            print(f"    * {s[0]} ({s[1]})")

    # 7. Exam Sessions
    print(f"\n[7] EXAM SESSIONS DETECTED IN GRADE CARDS:")
    rows = db.execute(text("SELECT id, name, exam_type, scheme FROM exam_sessions ORDER BY id")).fetchall()
    for r in rows:
        cnt = db.execute(text("SELECT COUNT(*) FROM grade_cards WHERE exam_session_id = :id"), {"id": r[0]}).scalar()
        print(f"  - Session [{r[0]}]: {r[1]:<25} | Type: {r[2] or 'N/A':<10} | Cards: {cnt:>3}")

    # 8. Grade Cards Breakdown
    print(f"\n[8] GRADE CARDS BY SEMESTER:")
    gc_stats = db.execute(text("""
        SELECT 
            semester, 
            COUNT(*) as count, 
            ROUND(MIN(sgpa), 2) as min_sgpa, 
            ROUND(AVG(sgpa), 2) as avg_sgpa, 
            ROUND(MAX(sgpa), 2) as max_sgpa,
            ROUND(AVG(cgpa), 2) as avg_cgpa
        FROM grade_cards 
        GROUP BY semester 
        ORDER BY semester
    """)).fetchall()
    print(f"  Sem | Count | Min SGPA | Avg SGPA | Max SGPA | Avg CGPA")
    print(f"  " + "-" * 55)
    for r in gc_stats:
        print(f"  S{r[0]:<2} | {r[1]:>5} | {r[2]:>8} | {r[3]:>8} | {r[4]:>8} | {r[5]:>8}")

    # 9. Semester Results (Grades)
    print(f"\n[9] SEMESTER RESULTS (SUBJECT GRADES):")
    grade_counts = db.execute(text("""
        SELECT grade, status, COUNT(*) as cnt
        FROM semester_results
        GROUP BY grade, status
        ORDER BY cnt DESC
    """)).fetchall()
    for r in grade_counts:
        print(f"  - Grade: {str(r[0]):<5} | Status: {str(r[1]):<6} : {r[2]:>4} entries")

    # 10. Course Offerings & Faculty
    print(f"\n[10] COURSE OFFERINGS & FACULTY:")
    fac_cnt = db.execute(text("SELECT COUNT(*) FROM faculty")).scalar()
    off_cnt = db.execute(text("SELECT COUNT(*) FROM course_offerings")).scalar()
    print(f"  - Total Faculty records : {fac_cnt}")
    print(f"  - Total Course Offerings: {off_cnt}")

    # 11. Data Integrity Checks
    print(f"\n[11] DATA INTEGRITY CHECKS:")
    orphaned_cards = db.execute(text("""
        SELECT COUNT(*) FROM grade_cards gc 
        LEFT JOIN students s ON gc.student_id = s.id 
        WHERE s.id IS NULL
    """)).scalar()
    print(f"  - Grade cards without student: {orphaned_cards}")

    orphaned_results = db.execute(text("""
        SELECT COUNT(*) FROM semester_results sr
        LEFT JOIN students s ON sr.student_id = s.id
        WHERE s.id IS NULL
    """)).scalar()
    print(f"  - Semester results without student: {orphaned_results}")

    orphaned_offerings = db.execute(text("""
        SELECT COUNT(*) FROM course_offerings co
        LEFT JOIN subjects s ON co.subject_id = s.id
        LEFT JOIN faculty f ON co.faculty_id = f.id
        WHERE s.id IS NULL OR f.id IS NULL
    """)).scalar()
    print(f"  - Offerings missing subject/faculty: {orphaned_offerings}")

    print("\n" + "=" * 65)
    print("                       VERIFICATION COMPLETE                    ")
    print("=" * 65)
    db.close()

if __name__ == "__main__":
    main()
